# -*- coding: utf-8 -*-
import openpyxl
import json
import re

def clean_key(s):
    if not s: return ''
    s = str(s).lower()
    s = s.replace('е', 'э').replace('й', 'и').replace('ё', 'о')
    s = re.sub(r'[\(\)\"\,\.\'\-\_]', ' ', s)
    s = re.sub(r'\b(ххк|хк|төүг|групп|корпораци|компани|llc|jsc|holding|холдинг|нэгдэл)\b', ' ', s)
    return re.sub(r'\s+', '', s)

def exact_key(s):
    if not s: return ''
    s = str(s).strip()
    s = re.sub(r'^(ххк|хк)\s+|\s+(ххк|хк)$', '', s, flags=re.IGNORECASE)
    s = re.sub(r'[\"\']', '', s)
    return s.strip()

print("1. Loading Top 100 data...")
wb100 = openpyxl.load_workbook('/Users/enkhbadraltsogbayar/Tender data/1_Монгол_Улсын_ТОП_100_ААН_2025_Бүрэн_мэдээлэл.xlsx', data_only=True)
ws100 = wb100.active

top100_by_key = {}
for r in range(2, ws100.max_row + 1):
    rank = ws100.cell(r, 1).value
    name = ws100.cell(r, 2).value
    reg = ws100.cell(r, 3).value
    web = ws100.cell(r, 4).value
    sh = ws100.cell(r, 5).value
    dr = ws100.cell(r, 6).value
    ubo = ws100.cell(r, 7).value
    vis = ws100.cell(r, 8).value
    mis = ws100.cell(r, 9).value
    
    entry = {
        'rank': rank, 'name': str(name).strip() if name else '',
        'reg_no': str(reg).strip() if reg else '',
        'web': str(web).strip() if web else '',
        'sh': str(sh).strip() if sh else '',
        'dr': str(dr).strip() if dr else '',
        'ubo': str(ubo).strip() if ubo else '',
        'vis': str(vis).strip() if vis else '',
        'mis': str(mis).strip() if mis else ''
    }
    k = clean_key(name)
    top100_by_key[k] = entry
    if reg:
        top100_by_key[f"reg_{reg}"] = entry

print(f"Top 100 indexed: {len(top100_by_key)} keys")

print("2. Loading LesOutputJson.json...")
with open('/Users/enkhbadraltsogbayar/Tender data/LesOutputJson.json', 'r', encoding='utf-8') as f:
    les_data = json.load(f)

les_by_key = {}
les_by_exact = {}
les_by_reg = {}

for row in les_data['Data']:
    reg = str(row[0]).strip()
    raw_name = str(row[1]).strip()
    les_by_reg[reg] = row
    
    ek = raw_name.lower()
    if ek not in les_by_exact:
        les_by_exact[ek] = row
        
    ck = clean_key(raw_name)
    if ck not in les_by_key:
        les_by_key[ck] = row

print(f"LesOutputJson indexed: {len(les_by_reg)} entities")

def parse_les_row(row):
    # row: [reg, name, date, form, type, sh, dr, ubo, act]
    reg = str(row[0])
    name = str(row[1])
    sh_list = []
    for sh in row[5]:
        if len(sh) >= 4 and (sh[2] or sh[3]):
            sh_list.append(f"{sh[2]} {sh[3]}".strip())
    dr_list = []
    for dr in row[6]:
        if len(dr) >= 4 and (dr[2] or dr[3]):
            dr_list.append(f"{dr[2]} {dr[3]}".strip())
    ubo_list = []
    for ub in row[7]:
        if len(ub) >= 3 and (ub[1] or ub[2]):
            ubo_list.append(f"{ub[1]} {ub[2]}".strip())
            
    return {
        'reg_no': reg,
        'official_name': name,
        'sh': ', '.join(dict.fromkeys(sh_list)),
        'dr': ', '.join(dict.fromkeys(dr_list)),
        'ubo': ', '.join(dict.fromkeys(ubo_list))
    }

print("3. Matching Top 300 companies...")
wb300 = openpyxl.load_workbook('/Users/enkhbadraltsogbayar/Downloads/Топ 300.xlsx', data_only=True)
ws300 = wb300.active

matched_results = []
unmatched_list = []

for r in range(3, ws300.max_row + 1):
    for c_type, n_col, name_col, em_col, web_col in [('ХХК', 2, 3, 4, 5), ('ХК', 7, 8, 9, 10)]:
        num = ws300.cell(r, n_col).value
        name = ws300.cell(r, name_col).value
        em = ws300.cell(r, em_col).value
        web = ws300.cell(r, web_col).value
        
        if not name:
            continue
            
        sname = str(name).strip()
        ck = clean_key(sname)
        
        # Check Top 100 first
        if ck in top100_by_key:
            t = top100_by_key[ck]
            matched_results.append({
                'source': 'Top 100 (100% verified)',
                'type': c_type, 'num': num, 'name': sname,
                'email': str(em or '').strip(), 'web': str(web or t['web']).strip(),
                'reg_no': t['reg_no'],
                'sh': t['sh'], 'dr': t['dr'], 'ubo': t['ubo'],
                'vis': t['vis'], 'mis': t['mis']
            })
            continue
            
        # Check LesOutputJson
        les_row = les_by_key.get(ck)
        if not les_row:
            raw_clean = sname.lower()
            les_row = les_by_exact.get(raw_clean)
            
        if les_row:
            p = parse_les_row(les_row)
            matched_results.append({
                'source': 'LesOutputJson',
                'type': c_type, 'num': num, 'name': sname,
                'email': str(em or '').strip(), 'web': str(web or '').strip(),
                'reg_no': p['reg_no'],
                'sh': p['sh'], 'dr': p['dr'], 'ubo': p['ubo'],
                'vis': '', 'mis': ''
            })
        else:
            unmatched_list.append({
                'type': c_type, 'num': num, 'name': sname,
                'email': str(em or '').strip(), 'web': str(web or '').strip()
            })

print(f"\nResults Summary:")
print(f"Total matched: {len(matched_results)} / 467 ({len(matched_results)/467*100:.1f}%)")
print(f"Unmatched: {len(unmatched_list)}")

if unmatched_list:
    print("\nSample unmatched companies (first 20):")
    for u in unmatched_list[:20]:
        print(f"[{u['type']} #{u['num']}] {u['name']} (web: {u['web']})")
